"""
Escáner de tickets de taxi -> registro de carreras (Fase 4, captura de datos)
TFM Pau Gavilán

Haz una foto del ticket en el momento, o sube varias fotos a la vez. Un modelo de
visión extrae fecha, hora, importe y km; tú los revisas/corriges y se añaden al
registro de carreras (Registro_carreras_TFM.xlsx).

Necesita una clave GRATUITA de Google AI Studio: https://aistudio.google.com/apikey

Coloca este archivo junto a Registro_carreras_TFM.xlsx y ejecuta:
    python -m pip install streamlit google-generativeai openpyxl pillow pandas
    python -m streamlit run escanear_tickets.py
"""

import io
import json
import os
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st

REGISTRO = "Registro_carreras_TFM.xlsx"
MODELO = "gemini-2.0-flash"   # modelo de visión (gratis). Puedes cambiarlo si hace falta.

PROMPT = (
    "Eres un asistente que lee tickets o recibos de taxi. Extrae los datos del ticket "
    "de la imagen y responde SOLO con un objeto JSON (sin texto alrededor) con estas "
    "claves, usando null cuando un dato no aparezca:\n"
    '{"fecha": "YYYY-MM-DD", "hora_recogida": "HH:MM", "importe_eur": number, '
    '"distancia_km": number}\n'
    "El importe es el total cobrado. Si el ticket usa coma decimal, conviértela a punto."
)


def _preparar_imagen(uploaded):
    """Devuelve (bytes, mime) reescalando si la foto es muy grande."""
    raw = uploaded.getvalue()
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(raw)).convert("RGB")
        img.thumbnail((1600, 1600))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=85)
        return buf.getvalue(), "image/jpeg"
    except Exception:
        return raw, uploaded.type or "image/jpeg"


def extraer(uploaded, api_key):
    import google.generativeai as genai
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel(MODELO)
    data, mime = _preparar_imagen(uploaded)
    resp = model.generate_content([PROMPT, {"mime_type": mime, "data": data}])
    txt = resp.text.strip()
    if txt.startswith("```"):
        txt = txt.strip("`")
        txt = txt[4:] if txt.lower().startswith("json") else txt
    return json.loads(txt.strip())


def parse_fecha(s):
    if not s:
        return s
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(s), fmt).date()
        except ValueError:
            pass
    return s


def añadir_al_registro(df):
    from openpyxl import load_workbook
    wb = load_workbook(REGISTRO)
    if "Registro" not in wb.sheetnames:
        raise ValueError("No encuentro la hoja 'Registro' en el Excel.")
    ws = wb["Registro"]
    cols = {c.value: c.column for c in ws[3] if c.value}   # cabecera en la fila 3
    fila = 4
    while ws.cell(fila, cols["Fecha"]).value not in (None, ""):
        fila += 1
    for _, r in df.iterrows():
        ws.cell(fila, cols["Fecha"], parse_fecha(r.get("fecha")))
        if "Hora recogida" in cols:
            ws.cell(fila, cols["Hora recogida"], r.get("hora_recogida"))
        ws.cell(fila, cols["Importe (€)"], r.get("importe_eur"))
        if "Distancia (km)" in cols and pd.notna(r.get("distancia_km")):
            ws.cell(fila, cols["Distancia (km)"], r.get("distancia_km"))
        fila += 1
    wb.save(REGISTRO)


def obtener_api_key():
    """Lee la clave de (1) .streamlit/secrets.toml, (2) variable de entorno; None si no hay."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY")


# =============================== UI ===============================
st.set_page_config(page_title="Escanear tickets de taxi", layout="centered")
st.title("Escanear tickets de taxi")
st.caption("Haz una foto del ticket o sube varias. Revisas los datos y se guardan en el registro.")

api_key = obtener_api_key()
if api_key:
    st.caption("Clave configurada correctamente.")
else:
    api_key = st.text_input("Clave de Google AI Studio (solo si no la has guardado)",
                            type="password", help="Gratis en aistudio.google.com/apikey")
    st.info("Consejo: guarda la clave una vez en `.streamlit/secrets.toml` y no tendrás "
            "que volver a pegarla.")

st.subheader("1. Añade los tickets")
foto = st.camera_input("Hacer una foto ahora")
subidas = st.file_uploader("...o subir varias fotos", type=["jpg", "jpeg", "png"],
                           accept_multiple_files=True)

imagenes = []
if foto is not None:
    imagenes.append(foto)
if subidas:
    imagenes.extend(subidas)

if imagenes and api_key and st.button(f"Leer {len(imagenes)} ticket(s)"):
    filas = []
    barra = st.progress(0.0)
    for i, img in enumerate(imagenes, 1):
        try:
            d = extraer(img, api_key)
            d["_archivo"] = getattr(img, "name", "foto")
            filas.append(d)
        except Exception as e:
            st.error(f"No pude leer {getattr(img, 'name', 'la foto')}: {e}")
        barra.progress(i / len(imagenes))
    if filas:
        st.session_state["tickets"] = pd.DataFrame(filas)

if "tickets" in st.session_state:
    st.subheader("2. Revisa y corrige")
    st.caption("Comprueba que los datos son correctos antes de guardar. Edita cualquier celda.")
    editado = st.data_editor(st.session_state["tickets"], use_container_width=True,
                             num_rows="dynamic")

    st.subheader("3. Guardar")
    if not Path(REGISTRO).exists():
        st.warning(f"No encuentro '{REGISTRO}' en esta carpeta.")
    elif st.button("Añadir al registro de carreras"):
        try:
            añadir_al_registro(editado)
            st.success(f"Añadidas {len(editado)} carrera(s) al registro.")
            del st.session_state["tickets"]
        except Exception as e:
            st.error(f"No pude guardar: {e}")
